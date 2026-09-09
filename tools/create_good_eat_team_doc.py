from pathlib import Path
from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\GameJam\GameJam")
OUT = ROOT / "output" / "好好吃饭队_游戏设计文档.docx"
ASSET_DIR = ROOT / "output" / "doc_assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

INSTRUCTION_IMG = Path(r"C:\Users\TU\AppData\Local\Temp\codex-clipboard-55ec4d71-2eb2-40ef-8dec-5763ac040ced.png")
TEAM_IMG = Path(r"C:\Users\TU\AppData\Local\Temp\codex-clipboard-e4ba371b-338d-4a5f-9567-62e510a55795.jpg")
TEAM_CROP = ASSET_DIR / "team_members_crop.jpg"
FLOW_IMG = ASSET_DIR / "core_loop_flow.png"

FONT_CJK = r"C:\Windows\Fonts\msyh.ttc"
NAVY = RGBColor(29, 49, 72)
BLUE = RGBColor(45, 103, 145)
ORANGE = RGBColor(208, 112, 52)
MUTED = RGBColor(100, 108, 116)
LIGHT_FILL = "F3F6F8"
BLUE_FILL = "EAF2F7"
ORANGE_FILL = "FFF2E9"
GRID = "D7E0E6"


def set_run_font(run, name="Microsoft YaHei", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if value and node is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not value and node is not None:
        p_pr.remove(node)


def set_picture_alt(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def style_paragraph(paragraph, size=11, color=None, bold=False, italic=False, align=None, before=0, after=6, line=1.25):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return paragraph


def add_para(doc, text="", size=11, color=None, bold=False, italic=False, align=None, before=0, after=6, line=1.25):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    style_paragraph(p, size=size, color=color, bold=bold, italic=italic, align=align, before=before, after=after, line=line)
    return p


def add_rich_para(doc, parts, before=0, after=6, line=1.25, align=None):
    p = doc.add_paragraph()
    for text, attrs in parts:
        r = p.add_run(text)
        set_run_font(r, **attrs)
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, size=16, color=BLUE, bold=True)
    elif level == 2:
        set_run_font(r, size=13, color=BLUE, bold=True)
    else:
        set_run_font(r, size=12, color=NAVY, bold=True)
    set_keep_with_next(p)
    return p


def add_caption(doc, text):
    return add_para(doc, text, size=9.5, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=10, line=1.1)


def add_callout(doc, label, text, fill=BLUE_FILL, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360], indent=120)
    set_table_borders(table, color=fill, size="0")
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.18
    r = p.add_run(label + "  ")
    set_run_font(r, size=11, color=accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=11, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.28 + 0.22 * level)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run_font(r, size=10.8, color=NAVY)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_run_font(r, size=10.8, color=NAVY)
    return p


def prepare_images():
    # The team screenshot contains a lot of phone UI below the member grid;
    # retaining the header and member area makes it useful in a design brief.
    img = Image.open(TEAM_IMG).convert("RGB")
    crop = img.crop((55, 165, 1025, 1120))
    crop.save(TEAM_CROP, quality=92)

    w, h = 1800, 640
    canvas = Image.new("RGB", (w, h), "#F7F9FA")
    draw = ImageDraw.Draw(canvas)
    try:
        font_big = ImageFont.truetype(FONT_CJK, 34)
        font_mid = ImageFont.truetype(FONT_CJK, 26)
        font_small = ImageFont.truetype(FONT_CJK, 22)
    except OSError:
        font_big = font_mid = font_small = ImageFont.load_default()
    boxes = [
        (60, 155, 335, 390, "选路线\n三选一奖励", "#EAF2F7"),
        (430, 155, 705, 390, "起手\n移动 / 上状态", "#FFF2E9"),
        (800, 155, 1075, 390, "放大\n目标与站位", "#EAF2F9"),
        (1170, 155, 1445, 390, "兑现\n击杀 / 返还", "#FFF2E9"),
    ]
    for x1, y1, x2, y2, label, fill in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=28, fill=fill, outline="#C9D6DE", width=3)
        lines = label.split("\n")
        ys = [y1 + 55, y1 + 112]
        for i, line in enumerate(lines):
            bbox = draw.textbbox((0, 0), line, font=font_big if i == 0 else font_mid)
            tw = bbox[2] - bbox[0]
            th = bbox[3] - bbox[1]
            draw.text(((x1 + x2 - tw) / 2, ys[i] - th / 2), line, fill="#1D3148", font=font_big if i == 0 else font_mid)
    for x in (355, 725, 1095):
        draw.line((x, 272, x + 55, 272), fill="#D07034", width=6)
        draw.polygon([(x + 55, 272), (x + 38, 261), (x + 38, 283)], fill="#D07034")
    draw.rounded_rectangle((420, 470, 1380, 570), radius=24, fill="#1D3148")
    note = "每回合最多回环一次：爽感有上限，决策才有价值"
    bbox = draw.textbbox((0, 0), note, font=font_small)
    draw.text(((w - (bbox[2] - bbox[0])) / 2, 495), note, fill="#FFFFFF", font=font_small)
    canvas.save(FLOW_IMG)


def set_document_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, color, before, after in [
        (1, 16, BLUE, 16, 8),
        (2, 13, BLUE, 12, 6),
        (3, 12, NAVY, 8, 4),
    ]:
        style = styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.8)


def set_page_layout(doc):
    section = doc.sections[0]
    # A4 is a named override for a Chinese submission document; all other
    # spacing and table tokens follow the selected narrative-proposal system.
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("好好吃饭队  ·  游戏设计文档")
    set_run_font(hr, size=9, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("GameJam / 肉鸽策略方案  ·  ")
    set_run_font(fr, size=9, color=MUTED)
    add_page_number(fp)


def add_cover(doc):
    add_para(doc, "GAMEJAM 设计文档", size=10, color=ORANGE, bold=True, before=18, after=20, line=1.0)
    p = add_para(doc, "好好吃饭队", size=28, color=NAVY, bold=True, after=4, line=1.0)
    set_keep_with_next(p)
    add_para(doc, "肉鸽策略感 × 连锁爽感的玩法方案", size=17, color=BLUE, after=18, line=1.15)
    add_para(doc, "基于现有 SRPG-光之阵项目的核心瓶颈回应", size=11.5, color=MUTED, after=22, line=1.15)

    add_callout(doc, "我们先给结论", "不要把爽感做成无脑加攻击。让玩家先搭一条短链，再用战棋的站位、目标顺序和行动力把它兑现；爽感来自“我做对了”，而不是来自随机数自己滚起来。", fill=ORANGE_FILL, accent=ORANGE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("小组成员")
    set_run_font(r, size=11, color=NAVY, bold=True)
    r = p.add_run("  陈泊森、宋瑞麒、李浩洪、贺文轩、胡浩鸿、孙爱琳")
    set_run_font(r, size=11, color=MUTED)

    pic = doc.add_picture(str(TEAM_CROP), width=Inches(4.55))
    set_picture_alt(pic, "小组成员截图", "好好吃饭队成员信息截图，显示六名成员头像和姓名。")
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "图 1  小组成员信息（按用户提供的队名整理为“好好吃饭队”）")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("文档用途")
    set_run_font(r, size=10, color=ORANGE, bold=True)
    r = p.add_run("  对应提交要求中的“核心瓶颈”部分，重点回答肉鸽构筑与战棋连锁反馈如何共存。")
    set_run_font(r, size=10, color=MUTED)


def add_team_and_context(doc):
    doc.add_page_break()
    add_heading(doc, "一、我们在做什么", 1)
    add_para(doc, "我们参考项目现有的《SRPG-光之阵》框架：地图是网格，角色有移动格数和行动回合，技能会受范围、障碍、阵营和资源消耗影响。项目本身已经有剑士、魔法师、弓箭手和医师四种主角定位，也有中毒、燃烧、出血、隐身、召唤和连击等技能或状态。", line=1.28)
    add_para(doc, "所以我们不打算另起一套完全独立的肉鸽战斗。更合适的做法，是在现有战棋底盘上加一层“本局构筑”：角色的职业和基础技能保持稳定，随机奖励负责改变这一局的行动方式。", line=1.28)
    add_callout(doc, "我们真正担心的点", "不是奖励数量不够，而是奖励拿到手以后不知道怎么用。如果奖励只是“攻击力 +10”，玩家很快就会选出答案；如果奖励能改变“先移动还是先上状态、先打谁、要不要继续追击”，它才会和战棋发生关系。", fill=BLUE_FILL, accent=BLUE)

    add_heading(doc, "附件信息怎么理解", 2)
    add_para(doc, "第一张附件是提交格式和内容建议，属于交付要求，不是游戏内规则；第二张附件用于确认小组成员。下面的方案正文以项目文档和用户提出的核心问题为准。", line=1.25)

    pic = doc.add_picture(str(INSTRUCTION_IMG), width=Inches(6.45))
    set_picture_alt(pic, "提交要求截图", "游戏设计文档提交要求截图，包含团队介绍、Benchmark、Core Loop 和核心瓶颈等内容建议。")
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "图 2  提交要求截图：团队介绍、Benchmark、Core Loop 和核心瓶颈")


def add_core_loop(doc):
    doc.add_page_break()
    add_heading(doc, "二、核心循环（Core Loop）", 1)
    add_para(doc, "我们的核心循环不是“打完一关，随机加数值，再打下一关”，而是让每一次奖励都回答一个问题：它能不能和我已经拥有的东西接上？", line=1.28)
    pic = doc.add_picture(str(FLOW_IMG), width=Inches(6.45))
    set_picture_alt(pic, "核心循环示意图", "核心循环示意图：选路线和奖励、起手、放大、兑现，并限制每回合最多回环一次。")
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "图 3  核心循环：奖励负责搭桥，战斗负责兑现")

    add_para(doc, "一局游戏中，玩家先在房间路线和三选一奖励里做决定，然后带着构筑进入战斗。战斗里，站位、目标顺序和行动力决定这条链能不能顺利走完；完成一次有效连锁后，玩家得到有限的再行动或资源返还，再决定是继续追击，还是保留优势结束回合。", line=1.28)
    add_heading(doc, "我们希望玩家记住的瞬间", 2)
    add_para(doc, "“我先用毒箭给右边的敌人挂状态，再走到安全格；击杀以后返还行动力，刚好够我去处理左边的精英。”", size=12, color=NAVY, italic=True, after=10, line=1.25)
    add_para(doc, "这类瞬间同时包含了策略和爽感：前半句是提前规划，后半句是即时兑现。", line=1.28)


def add_bottleneck(doc):
    doc.add_page_break()
    add_heading(doc, "三、核心瓶颈：肉鸽策略感，怎样和“左脚踩右脚”的爽感联动？", 1)
    add_para(doc, "这里的“左脚踩右脚”，我们理解成一种可控的自我联动：前一个动作产出状态、位置或资源，正好喂给下一个动作；下一个动作又给出一次有限的再行动。它应该让玩家感觉自己搭成了一套东西，但不能变成无限行动点。", line=1.28)

    add_callout(doc, "我们的判断", "采用“起手—放大—兑现”的搭桥式构筑。玩家决定链条怎么搭，战斗决定链条能不能成功；每回合最多回环一次，爽感有上限，策略才有价值。", fill=ORANGE_FILL, accent=ORANGE)

    add_heading(doc, "1. 奖励不是散装加成，而是链条里的一个位置", 2)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [1800, 3300, 4260], indent=120)
    set_table_borders(table)
    headers = ["位置", "它解决什么问题", "放进本项目的例子"]
    for i, text in enumerate(headers):
        cell = table.cell(0, i)
        shade_cell(cell, BLUE_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=10.2, color=NAVY, bold=True)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    mark_header_row(table.rows[0])
    rows = [
        ("起手", "先创造一个条件", "移动后标记、毒箭上中毒、火球施加燃烧"),
        ("放大", "满足条件后强化下一步", "攻击带状态目标时增伤、扩大范围或提高暴击"),
        ("兑现", "把结果转成下一次行动", "击杀后返还 1 点行动力、再移动 1 格或恢复 SP"),
    ]
    for row_data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(text)
            set_run_font(r, size=10, color=NAVY, bold=(i == 0))
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    add_caption(doc, "表 1  三种奖励位置：每张奖励只承担一个清楚的功能")

    add_heading(doc, "2. 用一个具体构筑把两种体验接起来", 2)
    add_para(doc, "以弓箭手的“毒箭连射”为例，我们可以把一条最小可玩链做成这样：", line=1.25)
    add_number(doc, "用毒箭连射给目标挂上中毒。它是起手，不要求玩家立刻打出最高伤害。")
    add_number(doc, "拿到“毒蚀瞄准”词条后，攻击中毒目标会获得额外伤害和暴击率。玩家需要判断先打普通敌人还是先处理精英。")
    add_number(doc, "如果成功击杀中毒目标，“收割步”返还 1 点行动力，并允许移动 1 格。这个返还只在本回合触发一次。")
    add_number(doc, "玩家可以继续追击，也可以停下来，避免为了多踩一步而走进敌人的包围。")
    add_para(doc, "这套设计的重点不在于数值很大，而在于每一步都能看见下一步的可能性。玩家会因为自己的顺序选择成功而觉得爽，而不是因为系统突然送了一次随机暴击。", line=1.28)

    add_heading(doc, "3. 为什么它不会把战棋策略冲掉", 2)
    add_bullet(doc, "奖励选择有机会成本：玩家是在起手、放大和兑现之间做取舍，而不是每次都拿同一种攻击力加成。")
    add_bullet(doc, "站位仍然重要：连锁需要目标距离、可通行格和安全撤退路线，不能只看角色面板。")
    add_bullet(doc, "目标顺序仍然重要：先打谁会决定状态是否浪费、返还是否能接上下一次攻击。")
    add_bullet(doc, "连锁有明确的断点：每回合最多一次行动力返还，整条链最多 2—3 个节点，避免无限滚雪球。")
    add_bullet(doc, "即使链条失败，第一段收益仍然有效，玩家不会因为一次 miss 就整回合失去体验。")

    add_heading(doc, "4. 爽感要在动作当下发生", 2)
    add_para(doc, "如果奖励只在战斗结算界面出现，玩家很难把奖励和自己的操作联系起来。我们的表现要求是：触发条件出现时目标高亮，命中后依次显示“中毒 → 增伤 → 暴击 → 击杀 → 返还行动力”，并立刻给出下一目标提示。", line=1.28)
    add_para(doc, "在实现上，真实伤害和行动力变化要放在战斗结算与事件触发阶段，而不是只改伤害飘字。项目文档已经指出，最终伤害应在 `GameBattleData.calculationHitResult` 一类的结算环节处理；肉鸽层负责记录本局的标签、触发次数和回滚数据。", size=10.5, color=MUTED, line=1.2)


def add_implementation(doc):
    doc.add_page_break()
    add_heading(doc, "四、第一版怎么落地", 1)
    add_para(doc, "第一版先验证“短链是否好玩”，不急着做完整随机地图和无限词条。我们会保留现有固定战斗场景，用三选一奖励加三种触发，尽快让玩家在一场战斗里感受到一次完整联动。", line=1.28)
    add_number(doc, "敌人死亡后生成三选一奖励，奖励明确标注“起手 / 放大 / 兑现”标签。")
    add_number(doc, "先实现“移动后首次命中强化”“命中带状态目标强化”“击杀返还一次行动力”三种效果。")
    add_number(doc, "奖励确认后立即写入肉鸽局状态；战斗中按回合清空触发次数，失败时清理本局技能、装备和属性增量。")
    add_number(doc, "战斗界面显示当前链条和本回合剩余返还次数，避免玩家不知道效果是否已经触发。")

    add_heading(doc, "我们会用什么标准判断它成功", 2)
    add_bullet(doc, "玩家在拿到奖励后，能用一句话说出它想接上的下一件东西。")
    add_bullet(doc, "一次普通行动能产生 2—3 个连续反馈，但不会出现无限行动。")
    add_bullet(doc, "链条是否成功取决于站位、目标顺序和资源，而不是单纯取决于角色面板数值。")
    add_bullet(doc, "连锁中断时仍然有局部收益，玩家会愿意下一回合继续尝试，而不是觉得构筑报废。")

    add_callout(doc, "最终方案", "把肉鸽放在“怎么搭构筑”，把爽感放在“这一回合能不能把构筑打出来”。起手、放大、兑现组成短链；短链只允许有限回环；位置、目标和行动力负责制造真正的决策。", fill=BLUE_FILL, accent=BLUE)

    add_heading(doc, "参考材料", 2)
    add_para(doc, "本方案参考项目内的《PROJECT_OVERVIEW.md》《ROGUELIKE_DESIGN_PLAN.md》和《CHARACTER_SKILL_EQUIPMENT_ANALYSIS.md》，沿用现有 SRPG 的角色、技能、状态、装备和战斗结算思路。本轮只整理设计文档，不修改游戏源码和数据。", size=10.2, color=MUTED, line=1.18)


def build():
    prepare_images()
    doc = Document()
    set_document_styles(doc)
    set_page_layout(doc)
    add_cover(doc)
    add_team_and_context(doc)
    add_core_loop(doc)
    add_bottleneck(doc)
    add_implementation(doc)
    props = doc.core_properties
    props.title = "好好吃饭队｜肉鸽策略感与连锁爽感玩法方案"
    props.subject = "基于 SRPG-光之阵项目的核心瓶颈回应"
    props.author = "好好吃饭队"
    props.keywords = "GameJam, 肉鸽, SRPG, 战棋, 构筑, 连锁爽感"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
